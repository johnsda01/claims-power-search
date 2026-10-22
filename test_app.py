import tempfile
import unittest
from pathlib import Path
from app import Library, normalized, extract, ocr_reading_text


class LookupTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.file = self.folder / 'Policy.txt'
        self.file.write_text('Power Failure\nFailure off the residence premises.\nBut a loss from a Peril Insured Against is covered.', encoding='utf-8')
        self.lib = Library(self.folder)
        self.ids = [f['id'] for f in self.lib.listing()['files']]

    def tearDown(self):
        self.temp.cleanup()

    def process(self):
        self.lib.worker(self.lib.selected(self.ids), False)
        self.assertEqual(self.lib.job['errors'], [])

    def test_conversion_preserves_source_and_context(self):
        original = self.file.read_bytes()
        self.process()
        self.assertEqual(original, self.file.read_bytes())
        text = (self.folder / 'claim-text/Policy.txt.txt').read_text(encoding='utf-8')
        self.assertIn('Peril Insured Against is covered.', text)
        self.assertEqual(len(self.lib.search(self.ids, '', ['food'])['results']), 1)

    def test_search_modes_and_topic_narrowing(self):
        self.process()
        self.assertEqual(self.lib.search(self.ids, 'power failure', [], 'phrase')['total'], 1)
        self.assertEqual(self.lib.search(self.ids, 'failure power', [], 'phrase')['total'], 0)
        self.assertEqual(self.lib.search(self.ids, 'failure power', [], 'all')['total'], 1)
        self.assertEqual(self.lib.search(self.ids, 'power', ['flood'])['total'], 0)
        self.assertEqual(self.lib.search(self.ids, '', ['food', 'flood'])['total'], 1)

    def test_reopen_and_folder_isolation(self):
        self.process()
        reopened = Library(self.folder)
        reopened.listing()
        self.assertEqual(reopened.search(self.ids, 'power', [])['total'], 1)
        other = self.folder / 'other'
        other.mkdir()
        (other / 'Policy.txt').write_text('Different policy. Flood is excluded.')
        reopened.listing(str(other))
        self.assertEqual(reopened.search(self.ids, 'power', [])['total'], 0)
        self.assertEqual(reopened.docs, {})

    def test_changed_source_requires_reprocess(self):
        self.process()
        self.file.write_text('Changed text. Tree debris removal.')
        result = self.lib.search(self.ids, 'power', [])
        self.assertEqual(result['total'], 0)
        self.assertEqual(result['skipped'], ['Policy.txt'])
        self.assertEqual(self.lib.listing()['files'][0]['status'], 'Not processed')

    def test_missing_source_removes_old_search_entry(self):
        self.process()
        self.file.unlink()
        self.lib.worker([self.file], False)
        self.assertNotIn('Policy.txt', self.lib.docs)
        self.assertTrue(self.lib.job['errors'])

    def test_invalid_selection_and_empty_search(self):
        with self.assertRaises(ValueError):
            self.lib.selected(['../elsewhere'])
        with self.assertRaises(ValueError):
            self.lib.search(self.ids, '', [])

    def test_dehyphenation_and_unicode(self):
        self.assertEqual(normalized('fail-\nure\u00a0of power'), 'failure of power')

    def test_word_paragraphs_and_tables(self):
        from docx import Document
        path = self.folder / 'endorsement.docx'
        doc = Document()
        doc.add_paragraph('Coverage extension')
        doc.add_table(rows=1, cols=1).cell(0,0).text = 'Food loss: $500'
        doc.save(path)
        result = extract(path, False, lambda _: None)
        self.assertIn('Coverage extension\nFood loss: $500', result[0]['text'])

    def test_ocr_rejoins_separate_clause_numbers(self):
        words = [{'text':'4.', 'x':10, 'y':100, 'width':10, 'height':10}, {'text':'Power Failure', 'x':40, 'y':100, 'width':70, 'height':10}]
        data = {'width':600, 'height':800, 'lines':[{'words':[w]} for w in words]}
        self.assertEqual(ocr_reading_text(data), '4. Power Failure')


if __name__ == '__main__':
    unittest.main()
