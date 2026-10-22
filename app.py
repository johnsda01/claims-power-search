"""Local document conversion and page-based policy lookup."""
from pathlib import Path
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse, parse_qs, quote
import argparse
import hashlib
import json
import mimetypes
import os
import re
import secrets
import subprocess
import sys
import tempfile
import threading
import unicodedata
import webbrowser

APP = Path(__file__).resolve().parent
APP_VERSION = (APP / 'VERSION').read_text(encoding='utf-8').strip()
from user_settings import UserSettings
SUPPORTED = {'.pdf', '.docx', '.txt', '.md'}
VERSION = 2
from page_highlights import GEOMETRY_LOCK as RENDER_LOCK
TOPICS = [
    ('trees', 'Trees & debris removal', ['tree', 'trees', 'debris removal', 'fallen limb', 'landscaping']),
    ('water', 'Sudden water discharge', ['accidental discharge', 'overflow of water', 'sudden and accidental', 'plumbing', 'household appliance']),
    ('leak', 'Repeated leaks & seepage', ['constant or repeated', 'seepage', 'leakage', '14 days', 'fourteen days']),
    ('flood', 'Flood & surface water', ['flood', 'surface water', 'storm surge', 'tidal water', 'waves']),
    ('backup', 'Sewer, drain & sump backup', ['sewer', 'sewers', 'drains', 'sump', 'backs up', 'backs-up']),
    ('ground', 'Groundwater & foundation seepage', ['below the surface', 'subsurface', 'hydrostatic', 'seeps', 'foundation']),
    ('roofwater', 'Rain through roof or openings', ['wind driven rain', 'wind-driven rain', 'opening in a roof', 'roof or wall', 'rain', 'roof leak']),
    ('freeze', 'Frozen pipes & appliances', ['freezing', 'frozen', 'maintain heat', 'drain all systems']),
    ('wind', 'Wind, hail & hurricanes', ['windstorm', 'hail', 'hurricane', 'wind or hail', 'windstorm or hail']),
    ('wear', 'Wear, deterioration & maintenance', ['wear and tear', 'wear, tear', 'deterioration', 'maintenance', 'mechanical breakdown', 'faulty', 'defective']),
    ('mold', 'Mold, fungi & rot', ['mold', 'fungi', 'fungus', 'wet rot', 'dry rot', 'bacteria']),
    ('food', 'Food loss & power failure', ['food', 'spoilage', 'spoiled', 'refrigerat', 'power failure', 'utility service', 'power interruption']),
    ('matching', 'Matching & undamaged property', ['matching', 'undamaged', 'cosmetic', 'aesthetic', 'uniform appearance']),
    ('settlement', 'Deductibles & payment limits', ['deductible', 'actual cash value', 'replacement cost', 'loss settlement', 'limit of liability', 'special limits']),
    ('duties', 'Notice, mitigation & claim duties', ['duties after loss', 'prompt notice', 'give notice', 'protect the property', 'neglect', 'proof of loss', 'reasonable emergency']),
    ('use', 'Loss of use & living expenses', ['additional living expense', 'loss of use', 'fair rental value', 'not fit to live', 'unfit for its normal use']),
]
COMMON = ['trees', 'wind', 'wear', 'mold', 'food', 'matching', 'settlement', 'duties', 'use', 'water']
WATER = ['water', 'leak', 'flood', 'backup', 'ground', 'roofwater', 'freeze']


def normalized(text):
    text = unicodedata.normalize('NFKC', text).lower()
    text = re.sub(r'(?<=\w)[-\u00ad]\s*\n\s*(?=\w)', '', text)
    return re.sub(r'\s+', ' ', text).strip()


def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def atomic_json(path, data):
    temp = path.with_suffix('.tmp')
    temp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    temp.replace(path)


def read_index(folder):
    try:
        value = json.loads((folder / 'claim-text' / 'index.json').read_text(encoding='utf-8'))
        if value.get('version') == VERSION and value.get('folder') == str(folder):
            return value['documents']
    except (OSError, ValueError, KeyError):
        pass
    return {}


def ocr_reading_text(data):
    words = [word for line in data['lines'] for word in line['words']]
    width, height = data['width'], data['height']
    body = [line for line in data['lines'] if line['words'] and .10 * height < line['words'][0]['y'] < .93 * height]
    spanning = sum(min(w['x'] for w in line['words']) < width*.43 and max(w['x']+w['width'] for w in line['words']) > width*.57 for line in body)
    columns = bool(body and spanning / len(body) < .25 and sum(w['x'] < width*.5 for w in words) > 30 and sum(w['x'] >= width*.5 for w in words) > 30)
    groups = [[], [], [], []] if columns else [[]]
    for word in words:
        center = word['y'] + word['height']/2
        group = 0 if not columns or center < height*.10 else 3 if center > height*.93 else 1 if word['x'] < width*.5 else 2
        groups[group].append(word)
    blocks = []
    for group in groups:
        lines = []
        for word in sorted(group, key=lambda w: (w['y']+w['height']/2, w['x'])):
            center = word['y']+word['height']/2
            if not lines or abs(center-lines[-1][0]) > word['height']*.45:
                lines.append((center, [word]))
            else:
                lines[-1][1].append(word)
        if lines:
            blocks.append('\n'.join(' '.join(w['text'] for w in sorted(line, key=lambda w:w['x'])) for _, line in lines))
    return '\n\n'.join(blocks)


def ocr_pages(path, numbers, progress):
    import pypdfium2
    result = {}
    with tempfile.TemporaryDirectory(prefix='claims-ocr-') as temp:
        temp = Path(temp)
        with pypdfium2.PdfDocument(path) as pdf:
            for number in numbers:
                page = pdf[number - 1]
                scale = min(2600 / max(page.get_size()), 4)
                bitmap = page.render(scale=scale)
                bitmap.to_pil().convert('RGB').save(temp / f'page-{number:05}.png')
                bitmap.close()
                page.close()
        progress(f'Reading {len(numbers)} scanned page(s) with Windows OCR…')
        shell = Path(os.environ.get('SystemRoot', 'C:/Windows')) / 'System32/WindowsPowerShell/v1.0/powershell.exe'
        run = subprocess.run([str(shell), '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass',
                              '-File', str(APP / 'ocr.ps1'), '-ImageFolder', str(temp)],
                             capture_output=True, text=True, timeout=max(120, len(numbers) * 30),
                             creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        if run.returncode:
            raise RuntimeError('Windows OCR failed. ' + (run.stderr.strip()[-600:] or 'Check that an OCR language is installed.'))
        for number in numbers:
            data = json.loads((temp / f'page-{number:05}.png.json').read_text(encoding='utf-8-sig'))
            result[number] = ocr_reading_text(data)
    return result


def extract(path, use_ocr, progress):
    if path.suffix.lower() == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('Password-protected PDF. Use an unlocked copy.')
        pages = []
        scan = []
        for number, page in enumerate(reader.pages, 1):
            progress(f'Extracting page {number} of {len(reader.pages)}…')
            text = page.extract_text() or ''
            method = 'PDF text'
            if len(re.sub(r'\W', '', text)) < 35:
                scan.append(number)
                method = 'Needs OCR'
            pages.append({'page': number, 'text': text, 'method': method})
        if scan and use_ocr:
            recognized = ocr_pages(path, scan, progress)
            for p in pages:
                if p['page'] in recognized:
                    p.update(text=recognized[p['page']], method='Windows OCR')
        return pages
    if path.suffix.lower() == '.docx':
        from docx import Document
        doc = Document(path)
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        blocks = []
        for child in doc.element.body:
            if child.tag.endswith('}p'):
                blocks.append(Paragraph(child, doc).text)
            elif child.tag.endswith('}tbl'):
                blocks.extend(' | '.join(c.text for c in row.cells) for row in Table(child, doc).rows)
        return [{'page': 1, 'text': '\n'.join(blocks), 'method': 'Word text; page numbers unavailable'}]
    raw = path.read_bytes()
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = raw.decode('utf-16') if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else raw.decode('cp1252')
    return [{'page': 1, 'text': text, 'method': 'Text file; page numbers unavailable'}]


class Library:
    def __init__(self, folder):
        self.lock = threading.RLock()
        self.folder = folder.resolve()
        self.docs = read_index(self.folder)
        self.allowed = {}
        self.job = {'running': False, 'message': 'Ready. Choose files and process them.', 'done': 0, 'total': 0, 'errors': [], 'warnings': []}
        self.cancel = threading.Event()

    def listing(self, folder=None):
        with self.lock:
            if folder:
                folder = Path(folder).expanduser().resolve()
                if not folder.is_dir():
                    raise ValueError('Choose an existing folder.')
                if self.job['running']:
                    raise ValueError('Wait for processing to finish before changing folders.')
                self.folder = folder
                self.docs = read_index(folder)
            files = []
            allowed = {}
            for path in sorted(self.folder.iterdir(), key=lambda p: p.name.lower()):
                if path.is_file() and path.suffix.lower() in SUPPORTED and not path.name.startswith('~$') and path.name.lower() != 'agents.md':
                    key = hashlib.sha256(path.name.encode()).hexdigest()[:20]
                    allowed[key] = path
                    item = self.docs.get(path.name)
                    stat = path.stat()
                    ready = bool(item and item.get('size') == stat.st_size and item.get('mtime') == stat.st_mtime_ns)
                    files.append({'id': key, 'name': path.name, 'size': stat.st_size,
                                  'status': ('Some pages need OCR' if any(p['method'] == 'Needs OCR' for p in item['pages']) else 'OCR used' if any(p['method'] == 'Windows OCR' for p in item['pages']) else 'Ready') if ready else 'Not processed',
                                  'pages': len(item['pages']) if ready else 0})
            self.allowed = allowed
            return {'folder': str(self.folder), 'output': str(self.folder / 'claim-text'), 'files': files}

    def process(self, ids, use_ocr):
        with self.lock:
            if self.job['running']:
                raise ValueError('Processing is already running.')
            paths = self.selected(ids)
            self.cancel.clear()
            self.job = {'running': True, 'message': 'Starting…', 'done': 0, 'total': len(paths), 'errors': [], 'warnings': []}
            thread = threading.Thread(target=self.worker, args=(paths, use_ocr), daemon=True)
            thread.start()
            return self.job.copy()

    def selected(self, ids):
        if not ids or not isinstance(ids, list):
            raise ValueError('Select at least one file.')
        if any(i not in self.allowed for i in ids):
            raise ValueError('The file list changed. Load the folder again.')
        return [self.allowed[i] for i in dict.fromkeys(ids)]

    def worker(self, paths, use_ocr):
        output = self.folder / 'claim-text'
        def progress(message):
            with self.lock:
                self.job['message'] = f'{path.name} · {message}'
        try:
            output.mkdir(exist_ok=True)
            for path in paths:
                if self.cancel.is_set():
                    break
                try:
                    signature = sha(path)
                    old = self.docs.get(path.name)
                    # Existing OCR may be reused, while an earlier skipped OCR pass may be retried.
                    if old and old.get('sha256') == signature and (not use_ocr or not any(p['method'] == 'Needs OCR' for p in old['pages'])) and (output / (path.name + '.txt')).exists():
                        progress('Text copy is current.')
                        stat = path.stat()
                        with self.lock:
                            old.update(size=stat.st_size, mtime=stat.st_mtime_ns)
                    else:
                        pages = extract(path, use_ocr, progress)
                        if sha(path) != signature:
                            raise ValueError('File changed while processing. Process it again.')
                        if not any(p['text'].strip() for p in pages):
                            raise ValueError('No readable text found. Enable OCR for scanned PDFs.')
                        warnings = [f"Page {p['page']}: {'OCR text; verify against the original.' if p['method'] == 'Windows OCR' else 'Little or no readable text; enable OCR.'}" for p in pages if p['method'] in ('Windows OCR', 'Needs OCR')]
                        text = f'SOURCE: {path.name}\nSHA256: {signature}\nSearch copy: verify exact wording in the original.\n'
                        for p in pages:
                            label = f"PDF PAGE {p['page']}" if path.suffix.lower() == '.pdf' else 'DOCUMENT TEXT (no page mapping)'
                            text += f"\n===== {label} | {p['method']} =====\n\n{p['text']}\n"
                        dest = output / (path.name + '.txt')
                        pending = dest.with_suffix('.txt.tmp')
                        pending.write_text(text, encoding='utf-8')
                        pending.replace(dest)
                        stat = path.stat()
                        with self.lock:
                            self.docs[path.name] = {'sha256': signature, 'size': stat.st_size, 'mtime': stat.st_mtime_ns,
                                                    'pages': pages, 'warnings': warnings}
                    with self.lock:
                        if self.docs[path.name].get('warnings'):
                            missing_text = any(p['method'] == 'Needs OCR' for p in self.docs[path.name]['pages'])
                            self.job['warnings'].append(path.name + (': Some pages have little or no text. Enable OCR and process again.' if missing_text else ': OCR used. Check wording you copy against the original.'))
                except Exception as e:
                    with self.lock:
                        self.docs.pop(path.name, None)
                        self.job['errors'].append(f'{path.name}: {e}')
                with self.lock:
                    self.job['done'] += 1
                    atomic_json(output / 'index.json', {'version': VERSION, 'folder': str(self.folder), 'documents': self.docs})
        except Exception as e:
            with self.lock:
                self.job['errors'].append(str(e))
        finally:
            with self.lock:
                self.job['running'] = False
                self.job['message'] = 'Stopped after the current file.' if self.cancel.is_set() else 'Processing complete.'

    def search(self, ids, query, topic_ids, mode='any'):
        with self.lock:
            paths = self.selected(ids)
            if self.job['running']:
                raise ValueError('Wait for processing to finish before searching.')
            catalog = {k: (label, words) for k, label, words in TOPICS}
            topics = [(key, *catalog[key]) for key in topic_ids if key in catalog]
            query = normalized(query)
            query_words = query.split()
            if not query and not topics:
                raise ValueError('Type a search or select a claim topic.')
            results, skipped = [], []
            for path in paths:
                doc = self.docs.get(path.name)
                if not doc or sha(path) != doc['sha256']:
                    skipped.append(path.name)
                    continue
                key = next(k for k, value in self.allowed.items() if value == path)
                for page in doc['pages']:
                    content = normalized(page['text'])
                    matches = [(label, [w for w in words if normalized(w) in content]) for _, label, words in topics]
                    hits = [(label, words) for label, words in matches if words]
                    query_hit = bool(query and (query in content if mode == 'phrase' else (all(w in content for w in query_words) if mode == 'all' else any(w in content for w in query_words))))
                    # A typed query narrows the selected topics. Topics themselves are OR searches.
                    if (query and not query_hit) or (topics and not hits):
                        continue
                    terms = list(dict.fromkeys(([query] if mode == 'phrase' else query_words) + [w for _, words in hits for w in words]))
                    results.append({'id': key, 'source': path.name, 'page': page['page'], 'text': page['text'],
                                    'method': page['method'], 'pdf': path.suffix.lower() == '.pdf',
                                    'topics': [label for label, _ in hits], 'terms': terms,
                                    'score': sum(content.count(normalized(t)) for t in terms),
                                    'textUrl': '/text/' + key, 'sourceUrl': '/source/' + key})
            results.sort(key=lambda r: (-r['score'], r['source'], r['page']))
            shown = results[:300]
            context = {item['id']: self.docs[self.allowed[item['id']].name]['pages'] for item in shown}
            return {'results': shown, 'total': len(results), 'skipped': skipped, 'context': context}


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def send(self, code, data, content_type='application/json; charset=utf-8'):
        if isinstance(data, (dict, list)):
            data = json.dumps(data, ensure_ascii=False).encode('utf-8')
        elif isinstance(data, str):
            data = data.encode('utf-8')
        self.send_response(code)
        self.send_header('Content-Type', content_type)
        self.send_header('Content-Length', str(len(data)))
        self.send_header('Cache-Control', 'no-store')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.send_header('Content-Security-Policy', "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; frame-src 'self'; object-src 'self'; frame-ancestors 'self'")
        self.end_headers()
        self.wfile.write(data)

    def check(self, mutation=False):
        host = self.headers.get('Host', '')
        valid = {f'127.0.0.1:{self.server.server_port}', f'localhost:{self.server.server_port}'}
        if host not in valid:
            raise ValueError('Only local connections are allowed.')
        origin = self.headers.get('Origin')
        if origin and origin not in {'http://' + value for value in valid}:
            raise ValueError('This request must come from the local app.')
        if mutation and self.headers.get('X-App-Token') != self.server.token:
            raise ValueError('Reload the app before continuing.')

    def do_GET(self):
        try:
            self.check()
            route = urlparse(self.path).path
            if route == '/':
                self.send(200, (APP / 'index.html').read_bytes(), 'text/html; charset=utf-8')
            elif route in ('/app.js', '/style.css'):
                self.send(200, (APP / route[1:]).read_bytes(), 'text/javascript; charset=utf-8' if route.endswith('.js') else 'text/css; charset=utf-8')
            elif route == '/api/init':
                preferences, warning = self.server.settings.load()
                self.send(200, {**self.server.library.listing(), 'token': self.server.token, 'topics': TOPICS, 'common': COMMON, 'water': WATER,
                                'version': APP_VERSION, 'settings': preferences, 'settingsWarning': warning, 'settingsPath': str(self.server.settings.path)})
            elif route == '/api/status':
                with self.server.library.lock:
                    self.send(200, self.server.library.job.copy())
            elif route.startswith('/page/'):
                import io
                import pypdfium2
                key = route.rsplit('/', 1)[-1]
                source = self.server.library.allowed.get(key)
                if not source or source.suffix.lower() != '.pdf':
                    raise ValueError('Choose a PDF from the current folder.')
                number = int(parse_qs(urlparse(self.path).query).get('page', ['1'])[0])
                with RENDER_LOCK, pypdfium2.PdfDocument(source) as pdf:
                    if not 1 <= number <= len(pdf):
                        raise ValueError('Page number is out of range.')
                    page = pdf[number - 1]
                    bitmap = page.render(scale=min(2200/max(page.get_size()), 4))
                    buffer = io.BytesIO()
                    bitmap.to_pil().convert('RGB').save(buffer, format='PNG')
                    bitmap.close()
                    page.close()
                self.send(200, buffer.getvalue(), 'image/png')
            elif route.startswith(('/source/', '/text/')):
                key = route.rsplit('/', 1)[-1]
                source = self.server.library.allowed.get(key)
                if not source:
                    raise ValueError('Load the folder again to open this file.')
                path = source if route.startswith('/source/') else source.parent / 'claim-text' / (source.name + '.txt')
                self.send(200, path.read_bytes(), 'application/pdf' if path.suffix.lower() == '.pdf' else 'text/plain; charset=utf-8' if path.suffix.lower() in ('.txt', '.md') else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            else:
                self.send(404, {'error': 'Not found.'})
        except Exception as e:
            self.send(400, {'error': str(e)})

    def do_POST(self):
        try:
            self.check(mutation=True)
            size = int(self.headers.get('Content-Length', '0'))
            if size > 1_000_000:
                raise ValueError('Request too large.')
            data = json.loads(self.rfile.read(size) or b'{}')
            route = urlparse(self.path).path
            lib = self.server.library
            if route == '/api/settings':
                result = {'settings': self.server.settings.update(data)}
            elif route == '/api/folder':
                result = lib.listing(data['folder'])
            elif route == '/api/browse':
                import tkinter as tk
                from tkinter import filedialog
                root = tk.Tk()
                root.withdraw()
                root.attributes('-topmost', True)
                try:
                    folder = filedialog.askdirectory(parent=root, initialdir=str(lib.folder), title='Choose a document folder')
                finally:
                    root.destroy()
                result = lib.listing(folder or None)
            elif route == '/api/process':
                result = lib.process(data.get('ids'), bool(data.get('ocr', True)))
            elif route == '/api/search':
                result = lib.search(data.get('ids'), data.get('query', ''), data.get('topics', []), data.get('mode', 'any'))
            elif route == '/api/highlights':
                from page_highlights import highlights
                source = lib.allowed.get(data.get('id'))
                if not source or source.suffix.lower() != '.pdf':
                    raise ValueError('Choose a PDF from the current folder.')
                segments = data.get('segments', [])
                if not isinstance(segments, list) or len(segments) > 300 or any(not isinstance(s, str) or len(s) > 5000 for s in segments):
                    raise ValueError('Invalid highlight text.')
                terms = data.get('terms')
                if terms is not None and (not isinstance(terms, list) or len(terms) > 300 or any(not isinstance(s, str) or len(s) > 5000 for s in terms)):
                    raise ValueError('Invalid highlight terms.')
                result = highlights(source, int(data.get('page', 1)), segments, terms, bool(data.get('focused', False)))
            elif route == '/api/stop':
                lib.cancel.set()
                result = {'ok': True}
            else:
                self.send(404, {'error': 'Not found.'})
                return
            self.send(200, result)
        except Exception as e:
            self.send(400, {'error': str(e)})


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--port', type=int, default=8777)
    parser.add_argument('--folder', default=str(APP.parent))
    parser.add_argument('--no-browser', action='store_true')
    args = parser.parse_args()
    try:
        server = ThreadingHTTPServer(('127.0.0.1', args.port), Handler)
    except OSError:
        from urllib.request import urlopen
        url = f'http://127.0.0.1:{args.port}'
        try:
            with urlopen(url + '/api/init', timeout=3) as response:
                existing = json.load(response)
            if 'topics' not in existing or 'token' not in existing:
                raise ValueError('Another program is using this port.')
        except Exception:
            raise SystemExit(f'Port {args.port} is in use. Close the other program or choose another port with --port.')
        if not args.no_browser:
            webbrowser.open(url)
        print('Claims Power Search is already running: ' + url)
        return
    server.library = Library(Path(args.folder))
    server.settings = UserSettings()
    server.token = secrets.token_urlsafe(32)
    url = f'http://127.0.0.1:{server.server_port}'
    print(url, flush=True)
    if not args.no_browser:
        webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == '__main__':
    main()
